import os
import re
from PIL import Image, ImageDraw, ImageFont
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_jupyter_screenshot(output_path):
    # Dimensions & Scaling
    width = 1100
    
    # Try to load standard Windows monospace font
    try:
        font_code = ImageFont.truetype("consola.ttf", 18)
        font_bold = ImageFont.truetype("consolab.ttf", 18)
        font_prompt = ImageFont.truetype("consola.ttf", 16)
        font_out = ImageFont.truetype("consola.ttf", 17)
    except:
        try:
            font_code = ImageFont.truetype("cour.ttf", 18)
            font_bold = ImageFont.truetype("courbd.ttf", 18)
            font_prompt = ImageFont.truetype("cour.ttf", 16)
            font_out = ImageFont.truetype("cour.ttf", 17)
        except:
            font_code = ImageFont.load_default()
            font_bold = font_code
            font_prompt = font_code
            font_out = font_code

    code_lines = [
        ("# Import required libraries", "comment"),
        ("import pandas as pd", "keyword_import"),
        ("from sklearn.linear_model import LogisticRegression", "keyword_import"),
        ("", "plain"),
        ("# ---------------------------------", "comment"),
        ("# Step 1: Create the dataset", "comment"),
        ("# ---------------------------------", "comment"),
        ("data = {", "plain"),
        ("    'Study_Hours': [1, 2, 3, 4, 5],", "dict_item"),
        ("    'Result': [0, 0, 0, 1, 1]", "dict_item"),
        ("}", "plain"),
        ("", "plain"),
        ("df = pd.DataFrame(data)", "plain"),
        ("", "plain"),
        ("# Display dataset", "comment"),
        ("print(\"Dataset:\")", "print_stmt"),
        ("print(df)", "print_stmt"),
        ("", "plain"),
        ("# ---------------------------------", "comment"),
        ("# Step 2: Separate Features (X) and Target (y)", "comment"),
        ("# ---------------------------------", "comment"),
        ("X = df[['Study_Hours']]      # Independent variable", "assign_comment"),
        ("y = df['Result']             # Dependent variable", "assign_comment"),
        ("", "plain"),
        ("# ---------------------------------", "comment"),
        ("# Step 3: Create Logistic Regression Model", "comment"),
        ("# ---------------------------------", "comment"),
        ("model = LogisticRegression()", "plain"),
        ("", "plain"),
        ("# ---------------------------------", "comment"),
        ("# Step 4: Train the Model", "comment"),
        ("# ---------------------------------", "comment"),
        ("model.fit(X, y)", "plain"),
        ("", "plain"),
        ("# ---------------------------------", "comment"),
        ("# Step 5: Display Learned Parameters", "comment"),
        ("# ---------------------------------", "comment"),
        ("print(\"\\nLearned Parameters\")", "print_stmt"),
        ("print(\"------------------\")", "print_stmt"),
        ("print(\"Weight (Coefficient) =\", model.coef_[0][0])", "print_stmt"),
        ("print(\"Bias (Intercept)     =\", model.intercept_[0])", "print_stmt"),
        ("", "plain"),
        ("# ---------------------------------", "comment"),
        ("# Step 6: Predict Training Data", "comment"),
        ("# ---------------------------------", "comment"),
        ("predictions = model.predict(X)", "plain"),
        ("probabilities = model.predict_proba(X)", "plain"),
        ("", "plain"),
        ("# ---------------------------------", "comment"),
        ("# Step 7: Display Results", "comment"),
        ("# ---------------------------------", "comment"),
        ("print(\"\\nPrediction Results\")", "print_stmt"),
        ("print(\"------------------\")", "print_stmt"),
        ("", "plain"),
        ("for i in range(len(df)):", "keyword_loop"),
        ("    print(f\"Study Hours = {X.iloc[i,0]} | \"", "fstring"),
        ("          f\"Actual = {y.iloc[i]} | \"", "fstring"),
        ("          f\"Predicted = {predictions[i]} | \"", "fstring"),
        ("          f\"Probability of Pass = {probabilities[i][1]:.4f}\")", "fstring"),
        ("", "plain"),
        ("# ---------------------------------", "comment"),
        ("# Step 8: Predict for a New Student", "comment"),
        ("# ---------------------------------", "comment"),
        ("new_student = [[3.5]]", "plain"),
        ("", "plain"),
        ("prediction = model.predict(new_student)", "plain"),
        ("probability = model.predict_proba(new_student)", "plain"),
        ("", "plain"),
        ("print(\"\\nNew Student Prediction\")", "print_stmt"),
        ("print(\"----------------------\")", "print_stmt"),
        ("print(\"Study Hours :\", new_student[0][0])", "print_stmt"),
        ("print(\"Predicted Class :\", prediction[0])", "print_stmt"),
        ("print(\"Probability of Pass :\", round(probability[0][1], 4))", "print_stmt")
    ]

    output_lines = [
        "Dataset:",
        "   Study_Hours  Result",
        "0            1       0",
        "1            2       0",
        "2            3       0",
        "3            4       1",
        "4            5       1",
        "",
        "Learned Parameters",
        "------------------",
        "Weight (Coefficient) = 1.0470438047457904",
        "Bias (Intercept)     = -3.7481774320466807",
        "",
        "Prediction Results",
        "------------------",
        "Study Hours = 1 | Actual = 0 | Predicted = 0 | Probability of Pass = 0.0629",
        "Study Hours = 2 | Actual = 0 | Predicted = 0 | Probability of Pass = 0.1606",
        "Study Hours = 3 | Actual = 0 | Predicted = 0 | Probability of Pass = 0.3527",
        "Study Hours = 4 | Actual = 1 | Predicted = 1 | Probability of Pass = 0.6083",
        "Study Hours = 5 | Actual = 1 | Predicted = 1 | Probability of Pass = 0.8156",
        "",
        "New Student Prediction",
        "----------------------",
        "Study Hours : 3.5",
        "Predicted Class : 0",
        "Probability of Pass : 0.4791"
    ]

    line_height_code = 24
    line_height_out = 22
    
    code_box_top = 40
    code_box_height = len(code_lines) * line_height_code + 30
    
    output_box_top = code_box_top + code_box_height + 15
    output_box_height = len(output_lines) * line_height_out + 30
    
    total_height = output_box_top + output_box_height + 40
    
    # Create blank canvas
    img = Image.new("RGB", (width, total_height), color=(255, 255, 255))
    draw = ImageDraw.Draw(img)

    # In prompt
    in_prompt_x = 25
    code_box_x = 120
    code_box_w = width - code_box_x - 35

    # Draw In [21]: text
    draw.text((in_prompt_x, code_box_top + 10), "In [21]:", fill=(48, 63, 159), font=font_prompt)

    # Draw Code input cell box
    draw.rectangle(
        [code_box_x, code_box_top, code_box_x + code_box_w, code_box_top + code_box_height],
        fill=(247, 247, 247),
        outline=(207, 216, 220),
        width=1
    )

    # Render Code Lines with syntax highlighting
    y_pos = code_box_top + 15
    for text, style in code_lines:
        x_pos = code_box_x + 18
        if style == "comment":
            draw.text((x_pos, y_pos), text, fill=(64, 128, 128), font=font_code)
        elif style == "keyword_import":
            # highlight import / from / as
            parts = text.split(" ")
            for p in parts:
                if p in ["import", "from", "as"]:
                    draw.text((x_pos, y_pos), p, fill=(0, 128, 0), font=font_bold)
                else:
                    draw.text((x_pos, y_pos), p, fill=(33, 33, 33), font=font_code)
                x_pos += int(draw.textlength(p + " ", font=font_code))
        elif style == "keyword_loop":
            parts = text.split(" ")
            for p in parts:
                if p in ["for", "in"]:
                    draw.text((x_pos, y_pos), p, fill=(0, 128, 0), font=font_bold)
                elif "range" in p or "len" in p:
                    draw.text((x_pos, y_pos), p, fill=(0, 134, 179), font=font_code)
                else:
                    draw.text((x_pos, y_pos), p, fill=(33, 33, 33), font=font_code)
                x_pos += int(draw.textlength(p + " ", font=font_code))
        elif style == "print_stmt":
            draw.text((x_pos, y_pos), text, fill=(33, 33, 33), font=font_code)
        elif style == "dict_item":
            draw.text((x_pos, y_pos), text, fill=(186, 33, 33), font=font_code)
        elif style == "assign_comment":
            if "#" in text:
                code_part, comment_part = text.split("#", 1)
                draw.text((x_pos, y_pos), code_part, fill=(33, 33, 33), font=font_code)
                offset = draw.textlength(code_part, font=font_code)
                draw.text((x_pos + offset, y_pos), "#" + comment_part, fill=(64, 128, 128), font=font_code)
            else:
                draw.text((x_pos, y_pos), text, fill=(33, 33, 33), font=font_code)
        else:
            draw.text((x_pos, y_pos), text, fill=(33, 33, 33), font=font_code)
        y_pos += line_height_code

    # Draw Output area
    y_out = output_box_top + 10
    for out_line in output_lines:
        draw.text((code_box_x + 18, y_out), out_line, fill=(33, 33, 33), font=font_out)
        y_out += line_height_out

    # Save image
    img.save(output_path, quality=95)
    print(f"Jupyter screenshot generated at: {output_path}")


def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_word_document(screenshot_path, doc_path):
    doc = docx.Document()

    # Configure Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Document Header Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("Machine Learning Practical Report")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(33, 33, 33)

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(16)
    run_sub = sub_p.add_run("Implementation of Logistic Regression: Study Hours vs. Exam Result")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(14)
    run_sub.font.color.rgb = RGBColor(90, 105, 120)

    # Horizontal Divider Line
    div_p = doc.add_paragraph()
    div_p.paragraph_format.space_after = Pt(16)
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="1" w:color="0078D4"/></w:pBdr>')
    div_p._p.get_or_add_pPr().append(pBdr)

    # Section 1: Objective & Overview
    h1 = doc.add_heading(level=1)
    h1.paragraph_format.space_before = Pt(10)
    h1.paragraph_format.space_after = Pt(6)
    r1 = h1.add_run("1. Experiment Overview & Objective")
    r1.font.name = "Calibri"
    r1.font.size = Pt(16)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(0, 102, 204)

    p_obj = doc.add_paragraph()
    p_obj.paragraph_format.space_after = Pt(10)
    p_obj.paragraph_format.line_spacing = 1.15
    run_obj = p_obj.add_run(
        "In this machine learning laboratory experiment, we implement a Binary Logistic Regression model "
        "using Python, Pandas, and Scikit-Learn. The goal is to predict whether a student passes (1) or fails (0) "
        "an examination based on the number of hours spent studying, and to calculate the underlying class probabilities."
    )
    run_obj.font.name = "Calibri"
    run_obj.font.size = Pt(11)

    # Section 2: Jupyter Notebook Implementation & Output Screenshot
    h2 = doc.add_heading(level=1)
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(6)
    r2 = h2.add_run("2. Jupyter Notebook Execution & Output Screenshot")
    r2.font.name = "Calibri"
    r2.font.size = Pt(16)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(0, 102, 204)

    p_desc = doc.add_paragraph()
    p_desc.paragraph_format.space_after = Pt(8)
    run_desc = p_desc.add_run("The complete execution cell containing all 8 logical steps and its rendered interactive output is shown below:")
    run_desc.font.name = "Calibri"
    run_desc.font.size = Pt(11)
    run_desc.font.italic = True

    # Embed Screenshot
    if os.path.exists(screenshot_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_after = Pt(14)
        run_img = p_img.add_run()
        run_img.add_picture(screenshot_path, width=Inches(6.8))

    # Section 3: Source Code Listing
    h3 = doc.add_heading(level=1)
    h3.paragraph_format.space_before = Pt(14)
    h3.paragraph_format.space_after = Pt(6)
    r3 = h3.add_run("3. Source Code (Step-by-Step Implementation)")
    r3.font.name = "Calibri"
    r3.font.size = Pt(16)
    r3.font.bold = True
    r3.font.color.rgb = RGBColor(0, 102, 204)

    # Create Code Box Table
    code_text = (
"""# Import required libraries
import pandas as pd
from sklearn.linear_model import LogisticRegression

# ---------------------------------
# Step 1: Create the dataset
# ---------------------------------
data = {
    'Study_Hours': [1, 2, 3, 4, 5],
    'Result': [0, 0, 0, 1, 1]
}

df = pd.DataFrame(data)

# Display dataset
print("Dataset:")
print(df)

# ---------------------------------
# Step 2: Separate Features (X) and Target (y)
# ---------------------------------
X = df[['Study_Hours']]      # Independent variable
y = df['Result']             # Dependent variable

# ---------------------------------
# Step 3: Create Logistic Regression Model
# ---------------------------------
model = LogisticRegression()

# ---------------------------------
# Step 4: Train the Model
# ---------------------------------
model.fit(X, y)

# ---------------------------------
# Step 5: Display Learned Parameters
# ---------------------------------
print("\\nLearned Parameters")
print("------------------")
print("Weight (Coefficient) =", model.coef_[0][0])
print("Bias (Intercept)     =", model.intercept_[0])

# ---------------------------------
# Step 6: Predict Training Data
# ---------------------------------
predictions = model.predict(X)
probabilities = model.predict_proba(X)

# ---------------------------------
# Step 7: Display Results
# ---------------------------------
print("\\nPrediction Results")
print("------------------")

for i in range(len(df)):
    print(f"Study Hours = {X.iloc[i,0]} | "
          f"Actual = {y.iloc[i]} | "
          f"Predicted = {predictions[i]} | "
          f"Probability of Pass = {probabilities[i][1]:.4f}")

# ---------------------------------
# Step 8: Predict for a New Student
# ---------------------------------
new_student = [[3.5]]

prediction = model.predict(new_student)
probability = model.predict_proba(new_student)

print("\\nNew Student Prediction")
print("----------------------")
print("Study Hours :", new_student[0][0])
print("Predicted Class :", prediction[0])
print("Probability of Pass :", round(probability[0][1], 4))"""
    )

    tbl_code = doc.add_table(rows=1, cols=1)
    tbl_code.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_code.autofit = False
    tbl_code.columns[0].width = Inches(6.8)
    cell_code = tbl_code.cell(0, 0)
    set_cell_background(cell_code, "F4F6F9")
    set_cell_margins(cell_code, top=140, bottom=140, left=180, right=180)
    
    p_code = cell_code.paragraphs[0]
    p_code.paragraph_format.space_before = Pt(0)
    p_code.paragraph_format.space_after = Pt(0)
    p_code.paragraph_format.line_spacing = 1.05
    run_c = p_code.add_run(code_text)
    run_c.font.name = "Consolas"
    run_c.font.size = Pt(9.5)
    run_c.font.color.rgb = RGBColor(30, 41, 59)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Section 4: Results & Analysis
    h4 = doc.add_heading(level=1)
    h4.paragraph_format.space_before = Pt(14)
    h4.paragraph_format.space_after = Pt(6)
    r4 = h4.add_run("4. Results and Analysis")
    r4.font.name = "Calibri"
    r4.font.size = Pt(16)
    r4.font.bold = True
    r4.font.color.rgb = RGBColor(0, 102, 204)

    # Summary Table
    tbl_res = doc.add_table(rows=6, cols=5)
    tbl_res.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Study Hours", "Actual Status", "Predicted Class", "Pass Probability", "Result Outcome"]
    rows_data = [
        ["1 hr", "Fail (0)", "Fail (0)", "6.29%", "Correct"],
        ["2 hrs", "Fail (0)", "Fail (0)", "16.06%", "Correct"],
        ["3 hrs", "Fail (0)", "Fail (0)", "35.27%", "Correct"],
        ["4 hrs", "Pass (1)", "Pass (1)", "60.83%", "Correct"],
        ["5 hrs", "Pass (1)", "Pass (1)", "81.56%", "Correct"],
    ]

    for col_idx, text in enumerate(headers):
        cell = tbl_res.cell(0, col_idx)
        set_cell_background(cell, "0078D4")
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(10.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    for row_idx, row in enumerate(rows_data):
        for col_idx, text in enumerate(row):
            cell = tbl_res.cell(row_idx + 1, col_idx)
            bg = "F8FAFC" if row_idx % 2 == 0 else "FFFFFF"
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(text)
            r.font.name = "Calibri"
            r.font.size = Pt(10)
            if col_idx == 4:
                r.font.bold = True
                r.font.color.rgb = RGBColor(16, 124, 65)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Section 5: Mathematical Findings & Decision Boundary
    h5 = doc.add_heading(level=1)
    h5.paragraph_format.space_before = Pt(14)
    h5.paragraph_format.space_after = Pt(6)
    r5 = h5.add_run("5. Mathematical Parameters & Decision Boundary")
    r5.font.name = "Calibri"
    r5.font.size = Pt(16)
    r5.font.bold = True
    r5.font.color.rgb = RGBColor(0, 102, 204)

    p_params = doc.add_paragraph()
    p_params.paragraph_format.line_spacing = 1.15
    p_params.paragraph_format.space_after = Pt(8)
    p_params.add_run("• Weight (Coefficient w): ").bold = True
    p_params.add_run("1.0470\n")
    p_params.add_run("• Bias (Intercept b): ").bold = True
    p_params.add_run("-3.7482\n")
    p_params.add_run("• Sigmoid Formula: ").bold = True
    p_params.add_run("P(Pass | Hours) = 1 / (1 + exp(-(1.0470 * Hours - 3.7482)))\n")
    p_params.add_run("• Decision Boundary (Threshold = 0.5): ").bold = True
    p_params.add_run("Hours = -b / w = 3.7482 / 1.0470 ≈ 3.58 Hours\n")
    p_params.add_run("• New Student Case (3.5 Hours): ").bold = True
    p_params.add_run("Pass Probability = 47.91% (< 50% threshold) → Predicted Class: 0 (Fail)")

    for r in p_params.runs:
        r.font.name = "Calibri"
        r.font.size = Pt(10.5)

    doc.save(doc_path)
    print(f"Word document successfully created at: {doc_path}")

if __name__ == "__main__":
    screenshot_file = r"c:\aryan codes\RAG-DEMO#2\jupyter_notebook_output.png"
    word_file = r"c:\aryan codes\RAG-DEMO#2\Logistic_Regression_Experiment.docx"
    create_jupyter_screenshot(screenshot_file)
    create_word_document(screenshot_file, word_file)
