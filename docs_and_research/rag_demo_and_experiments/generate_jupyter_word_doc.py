import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from PIL import Image, ImageDraw, ImageFont

# Set up output directories
output_dir = r"c:\aryan codes\RAG-DEMO#2"
img_dir = os.path.join(output_dir, "notebook_output_images")
os.makedirs(img_dir, exist_ok=True)

# Helper function to render high-resolution Jupyter output screenshots
def render_jupyter_output_image(text_lines, filename, has_out_prompt=False, out_num=None):
    scale = 2  # 2x scaling for high-DPI crisp screenshot in Word
    
    try:
        font = ImageFont.truetype("consola.ttf", 15 * scale)
        font_prompt = ImageFont.truetype("consola.ttf", 14 * scale)
    except:
        try:
            font = ImageFont.truetype("cour.ttf", 15 * scale)
            font_prompt = ImageFont.truetype("cour.ttf", 14 * scale)
        except:
            font = ImageFont.load_default()
            font_prompt = font

    line_height = 22 * scale
    padding_x = 24 * scale
    padding_y = 16 * scale
    
    # Calculate image dimensions
    max_line_len = max([len(line) for line in text_lines] + [10])
    img_width = max(900 * scale, (max_line_len * 10 + 60) * scale)
    img_height = (len(text_lines) * line_height) + (padding_y * 2)

    img = Image.new('RGB', (img_width, img_height), color='#FFFFFF')
    draw = ImageDraw.Draw(img)

    # Draw left subtle boundary or background
    # Subtle light gray box
    draw.rectangle([0, 0, img_width - 1, img_height - 1], outline='#E5E7EB', width=1*scale)

    # If Out prompt is needed (e.g. Out [4]:)
    current_y = padding_y
    for i, line in enumerate(text_lines):
        if i == 0 and has_out_prompt and out_num is not None:
            prompt_str = f"Out[{out_num}]: "
            draw.text((padding_x, current_y), prompt_str, fill='#D9534F', font=font_prompt)
            # offset text
            draw.text((padding_x + (len(prompt_str) * 9 * scale), current_y), line, fill='#212529', font=font)
        else:
            draw.text((padding_x, current_y), line, fill='#212529', font=font)
        current_y += line_height

    filepath = os.path.join(img_dir, filename)
    img.save(filepath, dpi=(150, 150))
    return filepath

# Generate screenshots for outputs of each cell that has output
out_files = {}

# Cell 2 Output
out_files[2] = render_jupyter_output_image([
    "Dataset:",
    "   Study_Hours  Result",
    "0            1       0",
    "1            2       0",
    "2            3       0",
    "3            4       1",
    "4            5       1"
], "cell2_output.png")

# Cell 4 Output (fit representation in Jupyter)
out_files[4] = render_jupyter_output_image([
    "LogisticRegression()"
], "cell4_output.png", has_out_prompt=True, out_num=4)

# Cell 5 Output
out_files[5] = render_jupyter_output_image([
    "Learned Parameters",
    "------------------",
    "Weight (Coefficient) = 1.0470438047457904",
    "Bias (Intercept)     = -3.7481774320466807"
], "cell5_output.png")

# Cell 6 Output
out_files[6] = render_jupyter_output_image([
    "Prediction Results",
    "------------------",
    "Study Hours = 1 | Actual = 0 | Predicted = 0 | Probability of Pass = 0.0629",
    "Study Hours = 2 | Actual = 0 | Predicted = 0 | Probability of Pass = 0.1606",
    "Study Hours = 3 | Actual = 0 | Predicted = 0 | Probability of Pass = 0.3527",
    "Study Hours = 4 | Actual = 1 | Predicted = 1 | Probability of Pass = 0.6083",
    "Study Hours = 5 | Actual = 1 | Predicted = 1 | Probability of Pass = 0.8156"
], "cell6_output.png")

# Cell 7 Output
out_files[7] = render_jupyter_output_image([
    "New Student Prediction",
    "----------------------",
    "Study Hours : 3.5",
    "Predicted Class : 0",
    "Probability of Pass : 0.4791"
], "cell7_output.png")

# Cell 8 (All-in-One) Output
out_files[8] = render_jupyter_output_image([
    "Dataset:",
    "   Study_Hours  Result",
    "0            1       0",
    "1            2       0",
    "2            3       0",
    "3            4       1",
    "4            5       1",
    "",
    "Learned Parameters",
    "------------------",
    "Weight (Coefficient) = 1.0470438047457904",
    "Bias (Intercept)     = -3.7481774320466807",
    "",
    "Prediction Results",
    "------------------",
    "Study Hours = 1 | Actual = 0 | Predicted = 0 | Probability of Pass = 0.0629",
    "Study Hours = 2 | Actual = 0 | Predicted = 0 | Probability of Pass = 0.1606",
    "Study Hours = 3 | Actual = 0 | Predicted = 0 | Probability of Pass = 0.3527",
    "Study Hours = 4 | Actual = 1 | Predicted = 1 | Probability of Pass = 0.6083",
    "Study Hours = 5 | Actual = 1 | Predicted = 1 | Probability of Pass = 0.8156",
    "",
    "New Student Prediction (3.5 hrs):",
    "Predicted Class: 0",
    "Probability of Pass: 0.4791"
], "cell8_all_output.png")

# ==========================================
# Build Word (.docx) Document
# ==========================================
doc = docx.Document()

# Page setup: 0.65 in margins for broad code cell display
for section in doc.sections:
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

# Helper for cell shading
def set_cell_background(cell, hex_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_borders(cell, top="#CFCFCF", bottom="#CFCFCF", left="#CFCFCF", right="#CFCFCF", sz="4"):
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{top.replace("#","")}"/>
            <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{left.replace("#","")}"/>
            <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{bottom.replace("#","")}"/>
            <w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{right.replace("#","")}"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)

def remove_cell_borders(cell):
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="none"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

# Title Header (Jupyter Notebook Bar)
title_p = doc.add_paragraph()
title_run = title_p.add_run("📓 logistic_regression.ipynb")
title_run.font.name = "Segoe UI"
title_run.font.size = Pt(18)
title_run.font.bold = True
title_run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
title_p.paragraph_format.space_after = Pt(4)

# Subtitle / Notebook Header banner
bar_p = doc.add_paragraph()
bar_run = bar_p.add_run("Jupyter Notebook • Python 3 (ipykernel)")
bar_run.font.name = "Segoe UI"
bar_run.font.size = Pt(10)
bar_run.font.italic = True
bar_run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
bar_p.paragraph_format.space_after = Pt(14)

# Markdown Cell 1
md1_p = doc.add_paragraph()
md1_h1 = md1_p.add_run("Logistic Regression - Pass/Fail Classification\n")
md1_h1.font.name = "Segoe UI"
md1_h1.font.size = Pt(15)
md1_h1.font.bold = True
md1_h1.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

md1_sub = md1_p.add_run("Simple implementation using pandas and scikit-learn.")
md1_sub.font.name = "Segoe UI"
md1_sub.font.size = Pt(10.5)
md1_sub.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
md1_p.paragraph_format.space_after = Pt(12)

# Helper function to add a Jupyter Code Cell & Output Screenshot
def add_jupyter_cell(doc, exec_num, code_lines, output_img_path=None):
    # Table layout with 2 columns: [In [x]: prompt] + [Code Block]
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Left cell: Prompt In [x]:
    cell_prompt = table.cell(0, 0)
    cell_prompt.width = Inches(0.9)
    remove_cell_borders(cell_prompt)
    set_cell_margins(cell_prompt, top=90, bottom=90, left=40, right=40)
    
    p_prompt = cell_prompt.paragraphs[0]
    p_prompt.paragraph_format.space_after = Pt(0)
    p_prompt.paragraph_format.space_before = Pt(0)
    p_prompt.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    
    r_prompt = p_prompt.add_run(f"In [{exec_num}]:")
    r_prompt.font.name = "Consolas"
    r_prompt.font.size = Pt(9.5)
    r_prompt.font.bold = True
    r_prompt.font.color.rgb = RGBColor(0x30, 0x3F, 0x9F)  # Jupyter Prompt Blue

    # Right cell: Code Box
    cell_code = table.cell(0, 1)
    cell_code.width = Inches(5.8)
    set_cell_background(cell_code, "F7F7F7")
    set_cell_borders(cell_code, top="#CFCFCF", bottom="#CFCFCF", left="#CFCFCF", right="#CFCFCF", sz="6")
    set_cell_margins(cell_code, top=80, bottom=80, left=120, right=120)

    p_code = cell_code.paragraphs[0]
    p_code.paragraph_format.space_after = Pt(0)
    p_code.paragraph_format.space_before = Pt(0)
    p_code.paragraph_format.line_spacing = 1.15

    for i, line in enumerate(code_lines):
        if i > 0:
            p_code = cell_code.add_paragraph()
            p_code.paragraph_format.space_after = Pt(0)
            p_code.paragraph_format.space_before = Pt(0)
            p_code.paragraph_format.line_spacing = 1.15
        
        r_code = p_code.add_run(line if line != "" else " ")
        r_code.font.name = "Consolas"
        r_code.font.size = Pt(9.5)
        
        # Simple syntax coloring for code text
        trimmed = line.strip()
        if trimmed.startswith("#"):
            r_code.font.color.rgb = RGBColor(0x6A, 0x73, 0x7D)  # Gray comment
            r_code.font.italic = True
        elif trimmed.startswith("import ") or trimmed.startswith("from "):
            r_code.font.color.rgb = RGBColor(0x00, 0x5C, 0xC5)  # Blue keyword
        elif trimmed.startswith("print(") or trimmed.startswith("print ("):
            r_code.font.color.rgb = RGBColor(0x6F, 0x42, 0xC1)  # Purple function
        else:
            r_code.font.color.rgb = RGBColor(0x24, 0x29, 0x2E)  # Dark charcoal code

    # Spacing after code cell
    post_code_p = doc.add_paragraph()
    post_code_p.paragraph_format.space_after = Pt(2)
    post_code_p.paragraph_format.space_before = Pt(2)

    # Output Screenshot Image (if any)
    if output_img_path and os.path.exists(output_img_path):
        # Add output table / container with Out prompt alignment
        out_table = doc.add_table(rows=1, cols=2)
        out_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        out_table.autofit = False

        c_out_prompt = out_table.cell(0, 0)
        c_out_prompt.width = Inches(0.9)
        remove_cell_borders(c_out_prompt)
        
        c_out_img = out_table.cell(0, 1)
        c_out_img.width = Inches(5.8)
        remove_cell_borders(c_out_img)
        set_cell_margins(c_out_img, top=40, bottom=40, left=0, right=0)
        
        p_img = c_out_img.paragraphs[0]
        p_img.paragraph_format.space_after = Pt(0)
        p_img.paragraph_format.space_before = Pt(0)
        
        # Add screenshot image
        run_img = p_img.add_run()
        run_img.add_picture(output_img_path, width=Inches(5.75))

        post_out_p = doc.add_paragraph()
        post_out_p.paragraph_format.space_after = Pt(10)
        post_out_p.paragraph_format.space_before = Pt(2)
    else:
        post_code_p.paragraph_format.space_after = Pt(10)

# -------------------------------------------------------------
# Add All Cells
# -------------------------------------------------------------

# Cell 1: Imports
add_jupyter_cell(doc, 1, [
    "# Step 0: Import Libraries",
    "import pandas as pd",
    "from sklearn.linear_model import LogisticRegression"
], None)

# Cell 2: Dataset Creation
add_jupyter_cell(doc, 2, [
    "# Step 1: Create the Dataset",
    "data = {",
    "    'Study_Hours': [1, 2, 3, 4, 5],",
    "    'Result': [0, 0, 0, 1, 1]",
    "}",
    "",
    "df = pd.DataFrame(data)",
    "print(\"Dataset:\")",
    "print(df)"
], out_files[2])

# Cell 3: Feature / Target Split
add_jupyter_cell(doc, 3, [
    "# Step 2: Separate Features (X) and Target (y)",
    "X = df[['Study_Hours']]      # Feature",
    "y = df['Result']             # Target label"
], None)

# Cell 4: Model Training
add_jupyter_cell(doc, 4, [
    "# Step 3: Initialize and Train Logistic Regression Model",
    "model = LogisticRegression()",
    "model.fit(X, y)"
], out_files[4])

# Cell 5: Learned Parameters
add_jupyter_cell(doc, 5, [
    "# Step 4: Display Learned Parameters",
    "print(\"Learned Parameters\")",
    "print(\"------------------\")",
    "print(\"Weight (Coefficient) =\", model.coef_[0][0])",
    "print(\"Bias (Intercept)     =\", model.intercept_[0])"
], out_files[5])

# Cell 6: Predictions
add_jupyter_cell(doc, 6, [
    "# Step 5: Predictions on Training Data",
    "predictions = model.predict(X)",
    "probabilities = model.predict_proba(X)",
    "",
    "print(\"Prediction Results\")",
    "print(\"------------------\")",
    "for i in range(len(df)):",
    "    print(f\"Study Hours = {X.iloc[i,0]} | \"",
    "          f\"Actual = {y.iloc[i]} | \"",
    "          f\"Predicted = {predictions[i]} | \"",
    "          f\"Probability of Pass = {probabilities[i][1]:.4f}\")"
], out_files[6])

# Cell 7: New Student Prediction
add_jupyter_cell(doc, 7, [
    "# Step 6: Prediction for a New Student (e.g. 3.5 Study Hours)",
    "new_student = [[3.5]]",
    "prediction = model.predict(new_student)",
    "probability = model.predict_proba(new_student)",
    "",
    "print(\"New Student Prediction\")",
    "print(\"----------------------\")",
    "print(\"Study Hours :\", new_student[0][0])",
    "print(\"Predicted Class :\", prediction[0])",
    "print(\"Probability of Pass :\", round(probability[0][1], 4))"
], out_files[7])

# Section Divider: All in one cell
md2_p = doc.add_paragraph()
md2_h2 = md2_p.add_run("### Complete All-in-One Cell")
md2_h2.font.name = "Segoe UI"
md2_h2.font.size = Pt(13)
md2_h2.font.bold = True
md2_h2.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
md2_p.paragraph_format.space_before = Pt(16)
md2_p.paragraph_format.space_after = Pt(8)

# Cell 8: All in one
add_jupyter_cell(doc, 8, [
    "# Complete Self-Contained Logistic Regression Script",
    "import pandas as pd",
    "from sklearn.linear_model import LogisticRegression",
    "",
    "# Dataset",
    "data = {'Study_Hours': [1, 2, 3, 4, 5], 'Result': [0, 0, 0, 1, 1]}",
    "df = pd.DataFrame(data)",
    "print(\"Dataset:\")",
    "print(df)",
    "",
    "# Split Features and Target",
    "X = df[['Study_Hours']]",
    "y = df['Result']",
    "",
    "# Model Training",
    "model = LogisticRegression()",
    "model.fit(X, y)",
    "",
    "# Model Parameters",
    "print(\"\\nLearned Parameters\")",
    "print(\"------------------\")",
    "print(\"Weight (Coefficient) =\", model.coef_[0][0])",
    "print(\"Bias (Intercept)     =\", model.intercept_[0])",
    "",
    "# Training Predictions",
    "predictions = model.predict(X)",
    "probabilities = model.predict_proba(X)",
    "",
    "print(\"\\nPrediction Results\")",
    "print(\"------------------\")",
    "for i in range(len(df)):",
    "    print(f\"Study Hours = {X.iloc[i,0]} | Actual = {y.iloc[i]} | Predicted = {predictions[i]} | Probability of Pass = {probabilities[i][1]:.4f}\")",
    "",
    "# New Inference",
    "new_student = [[3.5]]",
    "print(\"\\nNew Student Prediction (3.5 hrs):\")",
    "print(\"Predicted Class:\", model.predict(new_student)[0])",
    "print(\"Probability of Pass:\", round(model.predict_proba(new_student)[0][1], 4))"
], out_files[8])

output_docx_path = os.path.join(output_dir, "logistic_regression_jupyter_formatted.docx")
doc.save(output_docx_path)
print(f"Document created successfully at: {output_docx_path}")
